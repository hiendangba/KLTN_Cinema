from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUT = r"C:\hoctap\Study\KLTN\CinemaStar\cinema-microservices\chapter4_rewrite.docx"


def set_run_font(run, name="Times New Roman", size=13, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 13, bold=True)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    set_run_font(run, size=13, bold=False)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(13)

    add_heading(doc, "4.1. Triển khai hệ thống", level=1)
    add_body(
        doc,
        "Hệ thống CinemaStar được triển khai theo mô hình phân tách giữa lớp giao diện người dùng, lớp điều phối truy cập và các microservice backend. Cách tổ chức này giúp từng thành phần có thể được phát triển, đóng gói và vận hành độc lập, đồng thời vẫn đảm bảo sự phối hợp thống nhất thông qua API gateway và các cơ chế giao tiếp nội bộ.",
    )

    add_heading(doc, "4.1.1. Mô hình triển khai", level=2)
    add_body(
        doc,
        "Ở lớp ngoài cùng, người dùng tương tác với hệ thống thông qua ứng dụng web. Thành phần này đảm nhiệm việc hiển thị giao diện, tiếp nhận thao tác và gửi yêu cầu nghiệp vụ đến hệ thống ở phía sau.",
    )
    add_body(
        doc,
        "Lớp tiếp theo là API gateway, nơi tập trung việc xác thực, phân quyền và định tuyến yêu cầu đến đúng dịch vụ xử lý tương ứng. Nhờ đó, giao diện người dùng không cần biết chi tiết triển khai của từng service và hệ thống cũng dễ kiểm soát luồng truy cập hơn.",
    )
    add_body(
        doc,
        "Phần backend được chia thành nhiều microservice theo từng miền chức năng như xác thực, người dùng, phim, rạp chiếu, phòng chiếu, ghế, suất chiếu, đặt vé, thanh toán, đánh giá và gửi email. Cách chia này giúp giảm phụ thuộc chéo, hỗ trợ triển khai độc lập và thuận lợi khi mở rộng về sau.",
    )
    add_body(
        doc,
        "Tầng dữ liệu và hạ tầng bao gồm cơ sở dữ liệu, bộ nhớ đệm và các cơ chế xử lý bất đồng bộ. Cách tổ chức này giúp hệ thống vừa đảm bảo tính nhất quán dữ liệu, vừa nâng cao khả năng phản hồi đối với các tác vụ phát sinh trong quá trình vận hành.",
    )

    add_heading(doc, "4.1.2. Quy trình triển khai", level=2)
    add_body(
        doc,
        "Quy trình triển khai hệ thống được thực hiện theo từng bước rõ ràng để bảo đảm các thành phần sau khi đóng gói có thể khởi động và phối hợp ổn định với nhau.",
    )
    add_body(
        doc,
        "Trước hết, mã nguồn được chuẩn bị thành hai khối chính là frontend và backend. Frontend được xây dựng thành ứng dụng web riêng, còn backend được tách thành các service độc lập đi kèm với các thành phần hạ tầng như cơ sở dữ liệu, cache, message queue và lớp điều phối truy cập.",
    )
    add_body(
        doc,
        "Sau giai đoạn đóng gói, các service được cấu hình để giao tiếp thông qua API gateway và các giao diện nội bộ. Cách tiếp cận này cho phép từng thành phần được triển khai đúng vai trò kỹ thuật của mình nhưng vẫn bảo đảm luồng xử lý tổng thể của hệ thống.",
    )
    add_body(
        doc,
        "Khi đưa lên môi trường triển khai, hệ thống được kiểm tra trạng thái sẵn sàng, khả năng kết nối giữa các service và khả năng phản hồi của các luồng nghiệp vụ chính. Việc kiểm tra tập trung vào mức độ ổn định của hệ thống sau khi triển khai thay vì chỉ xác nhận từng thành phần hoạt động riêng lẻ.",
    )

    add_heading(doc, "4.2. Kiểm thử và đánh giá hệ thống", level=1)
    add_body(
        doc,
        "Quá trình kiểm thử của CinemaStar được thực hiện ở hai mức chính: unit test trong mã nguồn và kiểm thử API tự động. Cách tiếp cận này giúp vừa phát hiện sớm lỗi logic ở từng thành phần, vừa xác minh khả năng vận hành của các luồng nghiệp vụ quan trọng trên môi trường chạy thật.",
    )

    add_heading(doc, "4.2.1. Mục tiêu kiểm thử", level=2)
    add_body(
        doc,
        "Mục tiêu của kiểm thử là xác minh tính đúng đắn của các rule nghiệp vụ ở mức đơn vị, đồng thời đánh giá khả năng hoạt động của các API chính trên môi trường production. Phạm vi kiểm thử tập trung vào các chức năng quan trọng như xác thực người dùng, tra cứu phim, tra cứu suất chiếu, xem sơ đồ ghế, đặt vé, thanh toán, đánh giá, báo cáo doanh thu và xuất báo cáo Excel.",
    )

    add_heading(doc, "4.2.2. Công cụ và môi trường kiểm thử", level=2)
    add_body(
        doc,
        "Ở mức unit test, hệ thống sử dụng JUnit và Mockito để kiểm tra độc lập từng service, controller và các thành phần nội bộ. Ở mức API, Postman Collection kết hợp Newman CLI được dùng để chạy tự động các request và xuất báo cáo HTML/JUnit. Các request được gửi tới API gateway production tại địa chỉ https://cinema-api.duckdns.org.",
    )

    add_heading(doc, "4.2.3. Kiểm thử đơn vị (Unit Test)", level=2)
    add_body(
        doc,
        "Trong quá trình phát triển CinemaStar, nhóm xây dựng bộ unit test cho các lớp xử lý cốt lõi nhằm kiểm tra sớm từng nhánh nghiệp vụ trước khi tích hợp vào hệ thống hoàn chỉnh. Các test này giúp xác minh dữ liệu đầu vào, kết quả trả về, xử lý lỗi và các điều kiện ràng buộc ở từng module riêng lẻ.",
    )
    add_body(
        doc,
        "Nhóm unit test tiêu biểu cho xác thực và người dùng gồm RequestAuthUtilsTest, JwtAuthenticationFilterTest, UserServiceImplTokenFlowTest, UserServiceImplPhoneLookupTest, CustomerRankServiceImplTest và LoyaltyPointsSyncServiceTest. Nhóm test cho phim, suất chiếu và ghế gồm FilmServiceImplTest, FilmControllerSearchIntegrationTest, ActorServiceImplTest, TypeServiceImplTest, ShowTimeServiceImplTest và SeatSuggestionServiceImplTest.",
    )
    add_body(
        doc,
        "Nhóm unit test cho đặt vé, thanh toán và các module hỗ trợ gồm BookingServiceImplTest, BookingInternalGrpcServiceTest, PaymentSessionServiceImplTest, PromotionServiceImplTest, PromotionEngineTest, ReviewServiceImplTest, UploadServiceImplTest, HallServiceImplTest và CinemaServiceImplTest. Nhìn chung, bộ unit test đóng vai trò lớp kiểm tra sơ bộ, giúp phát hiện sớm lỗi logic ngay trong mã nguồn và tạo nền tảng tin cậy cho các bước kiểm thử tích hợp tiếp theo.",
    )

    add_heading(doc, "4.2.4. Kiểm thử chức năng API bằng Postman/Newman", level=2)
    add_body(
        doc,
        "Postman Collection được sử dụng để nhóm các request theo từng luồng nghiệp vụ, trong khi Newman CLI đảm nhiệm việc chạy tự động collection, tổng hợp kết quả và xuất báo cáo HTML/JUnit. Cách làm này giúp quá trình kiểm thử API có thể lặp lại nhiều lần trên API gateway production mà không phải thao tác thủ công từng request.",
    )
    add_body(
        doc,
        "Các luồng được kiểm thử tập trung vào những chức năng người dùng sử dụng thường xuyên như đăng nhập, tra cứu phim, xem suất chiếu, xem sơ đồ ghế, đặt vé và thanh toán. Việc kiểm thử theo luồng giúp đánh giá được mức độ ổn định của hệ thống trong điều kiện sử dụng thực tế.",
    )

    add_heading(doc, "4.2.5. Kết quả kiểm thử", level=2)
    add_body(
        doc,
        "Sau khi chạy collection bằng Newman trên môi trường production, hệ thống thực thi 19 request với tổng cộng 57 assertion. Kết quả cho thấy toàn bộ assertion đều thành công và không có test thất bại.",
    )
    add_body(
        doc,
        "Ngoài kết quả tổng quan, báo cáo Newman còn cung cấp thông tin chi tiết cho từng request. Ví dụ, request POST /api/films/search được sử dụng để kiểm tra chức năng tìm kiếm phim, trả về mã trạng thái 200 OK và cho thấy luồng kiểm thử hoạt động ổn định trên môi trường mục tiêu.",
    )

    doc.save(OUT)


if __name__ == "__main__":
    main()
