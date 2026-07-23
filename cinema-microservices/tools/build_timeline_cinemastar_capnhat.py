from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = r"C:\hoctap\Study\KLTN\File Báo cáo\Timeline_CinemaStar_Nhap_capnhat.docx"


TIMELINE_ROWS = [
    (
        "1",
        "24/02/2026 - 01/03/2026",
        "Khảo sát yêu cầu, phân tích quy trình bán vé rạp phim, xác định actor, use case và phạm vi chức năng cho khách hàng, nhân viên và quản trị viên.",
        "Khảo sát yêu cầu, Use case tổng thể",
        "Hình thành bộ yêu cầu chức năng, sơ đồ use case và định hướng nghiệp vụ cho toàn hệ thống.",
    ),
    (
        "2",
        "02/03/2026 - 08/03/2026",
        "Thiết kế kiến trúc microservices, phân tách domain, xác định service giao tiếp qua API gateway, gRPC và các thành phần hạ tầng dùng chung.",
        "Kiến trúc hệ thống, API gateway, common-lib",
        "Hoàn thiện kiến trúc tổng thể và khung giao tiếp giữa các service.",
    ),
    (
        "3",
        "09/03/2026 - 15/03/2026",
        "Xây dựng chức năng đăng nhập, refresh token, phân quyền theo role và cơ chế xác thực JWT; đồng thời chuẩn bị flow Google OAuth cho đăng nhập ngoài hệ thống.",
        "Identity service, Auth, Google OAuth",
        "Hoàn thiện nền tảng xác thực và phân quyền cho toàn bộ hệ thống.",
    ),
    (
        "4",
        "16/03/2026 - 22/03/2026",
        "Phát triển các nghiệp vụ người dùng như quản lý hồ sơ, customer profile, tìm kiếm theo số điện thoại, xếp hạng khách hàng và đồng bộ điểm tích lũy.",
        "User service, Customer rank, Loyalty points",
        "Hoàn thiện module quản lý người dùng và loyalty workflow.",
    ),
    (
        "5",
        "23/03/2026 - 29/03/2026",
        "Thiết kế và cài đặt dữ liệu rạp, phòng chiếu, sơ đồ ghế, cấu hình hall và seat layout phục vụ cho bước tạo suất chiếu và bán vé.",
        "Cinema service, Hall service, Seat service",
        "Xây dựng xong nền dữ liệu vật lý của rạp chiếu phim.",
    ),
    (
        "6",
        "30/03/2026 - 05/04/2026",
        "Triển khai quản lý phim, diễn viên, thể loại, thông tin hiển thị và chức năng tìm kiếm phim phục vụ cả người dùng lẫn quản trị viên.",
        "Film service, Actor, Type, Search film",
        "Hoàn thiện module phim và chức năng tra cứu danh mục phim.",
    ),
    (
        "7",
        "06/04/2026 - 12/04/2026",
        "Phát triển chức năng tạo suất chiếu, tìm kiếm suất chiếu theo nhiều điều kiện, áp dụng chính sách giá vé và gợi ý ghế phù hợp theo trạng thái thực tế.",
        "Showtime service, Pricing policy, Seat suggestion",
        "Hoàn thiện nhóm chức năng liên quan đến lịch chiếu và dữ liệu hiển thị trước khi đặt vé.",
    ),
    (
        "8",
        "13/04/2026 - 19/04/2026",
        "Cài đặt luồng giữ ghế, tạo booking, xử lý combo đi kèm và kiểm soát trạng thái ghế trong quá trình người dùng thao tác đặt vé.",
        "Booking service, Product service, Seat hold",
        "Hình thành luồng đặt vé cơ bản và cơ chế giữ ghế tạm thời.",
    ),
    (
        "9",
        "20/04/2026 - 26/04/2026",
        "Phát triển module thanh toán, session thanh toán, khuyến mãi, áp mã giảm giá và đồng bộ dữ liệu doanh thu phát sinh sau giao dịch.",
        "Payment service, Promotion, Revenue",
        "Hoàn thiện nghiệp vụ thanh toán và khung xử lý doanh thu.",
    ),
    (
        "10",
        "27/04/2026 - 03/05/2026",
        "Bổ sung chức năng đánh giá phim, kiểm tra eligibility review, upload hình ảnh/video và các tác vụ hỗ trợ hiển thị nội dung đánh giá.",
        "Review service, Upload service",
        "Mở rộng hệ thống với chức năng review và upload nội dung đa phương tiện.",
    ),
    (
        "11",
        "04/05/2026 - 10/05/2026",
        "Xây dựng giao diện frontend cho các màn hình chính như đăng nhập, trang chủ, danh sách phim, chi tiết phim, lịch chiếu và sơ đồ ghế.",
        "Frontend customer pages",
        "Hoàn thiện khung giao diện phía người dùng cuối cho các luồng chính.",
    ),
    (
        "12",
        "11/05/2026 - 17/05/2026",
        "Phát triển giao diện quản trị và bán vé tại quầy, bao gồm quản lý phim, rạp, phòng, suất chiếu, khách hàng và thao tác bán vé cho staff.",
        "Admin UI, Staff-sell UI",
        "Hoàn thiện nhóm chức năng frontend cho quản trị viên và nhân viên bán vé.",
    ),
    (
        "13",
        "18/05/2026 - 24/05/2026",
        "Tích hợp các service qua gRPC và API gateway, chuẩn hóa contract response, đồng bộ auth header và hoàn thiện các luồng gọi nội bộ giữa service.",
        "gRPC nội bộ, Envoy gateway, service integration",
        "Đồng nhất cơ chế giao tiếp và giảm phụ thuộc chéo giữa các service.",
    ),
    (
        "14",
        "25/05/2026 - 31/05/2026",
        "Bổ sung unit test cho các lớp service, controller, scheduler và gRPC; đồng thời xây dựng Postman Collection để kiểm thử các API chính theo từng luồng nghiệp vụ.",
        "Unit test, API test, Postman/Newman",
        "Hoàn thiện lớp kiểm thử đơn vị và kiểm thử API tự động cho hệ thống.",
    ),
    (
        "15",
        "01/06/2026 - 07/06/2026",
        "Triển khai hệ thống trên môi trường Docker Compose, cấu hình gateway, database, cache, message queue và kiểm tra khả năng chạy đồng thời của các service.",
        "Triển khai hệ thống, Docker Compose, Redis, RabbitMQ",
        "Hệ thống có thể khởi động và vận hành theo mô hình microservices trên môi trường tích hợp.",
    ),
    (
        "16",
        "08/06/2026 - 14/06/2026",
        "Tối ưu các luồng nghiệp vụ, sửa lỗi sau kiểm thử, bổ sung export Excel, báo cáo doanh thu, log trace và hoàn thiện các chức năng còn thiếu.",
        "Fix bug, Report, Export Excel, Trace log",
        "Nâng cao độ ổn định, tính quan sát và chất lượng đầu ra của hệ thống.",
    ),
    (
        "17",
        "15/06/2026 - 02/07/2026",
        "Tổng hợp kết quả thực hiện, hoàn thiện báo cáo, chuẩn bị slide, rà soát mã nguồn, kiểm tra demo end-to-end và chỉnh sửa lần cuối trước khi nộp.",
        "Báo cáo, Slide, Demo, Tổng rà soát",
        "Hoàn thiện hồ sơ đồ án và chuẩn bị cho giai đoạn bảo vệ tốt nghiệp.",
    ),
]


def set_run_font(run, size=12, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def format_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=False):
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    if first_indent:
        paragraph.paragraph_format.first_line_indent = Cm(0.75)


def set_cell_text(cell, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=11):
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        format_paragraph(p, align=align)
        for r in p.runs:
            set_run_font(r, size=size, bold=bold)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.54)
    sec.right_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("TIMELINE THỰC HIỆN CÔNG VIỆC CINEMASTAR")
    set_run_font(r, size=16, bold=True)

    for text in [
        "Bản nháp riêng trình bày theo dạng đề cương tiến độ",
        "Tiến độ được chia theo từng giai đoạn thực hiện module và các đầu ra chính của hệ thống",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        set_run_font(r, size=12)

    intro = doc.add_paragraph()
    format_paragraph(intro, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=True)
    r = intro.add_run(
        "Bảng tiến độ dưới đây được trình bày theo hướng giống đề cương thực hiện khóa luận, "
        "trong đó mỗi mốc thời gian thể hiện rõ module hoặc nhóm chức năng được triển khai, "
        "nội dung công việc chính và kết quả đầu ra dự kiến. Các mốc thời gian được sắp xếp "
        "theo trình tự phát triển hợp lý của một hệ thống microservices bán vé rạp phim."
    )
    set_run_font(r, size=12)

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    widths = [Cm(1.1), Cm(3.4), Cm(6.2), Cm(4.5), Cm(4.8)]
    headers = ["STT", "Thời gian", "Nội dung thực hiện", "Module / thành phần", "Kết quả / ghi chú"]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_text(cell, header, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=11)

    for row_data in TIMELINE_ROWS:
        row = table.add_row()
        for i, value in enumerate(row_data):
            row.cells[i].width = widths[i]
            align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 1) else WD_ALIGN_PARAGRAPH.JUSTIFY
            if i == 3:
                align = WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[i], value, align=align, size=11)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(0)
    r = note.add_run(
        "Ghi chú: Nếu cần rút gọn khi chèn vào báo cáo chính, có thể giữ lại cột thời gian, "
        "nội dung thực hiện và module / thành phần; cột kết quả / ghi chú dùng để tăng tính "
        "thuyết minh khi trình bày tiến độ."
    )
    set_run_font(r, size=11, bold=False)

    doc.save(OUT)


if __name__ == "__main__":
    main()
