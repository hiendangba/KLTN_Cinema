import os

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


def set_text(paragraph, text):
    paragraph.text = text


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style is not None:
        new_para.style = style
    if text:
        new_para.text = text
    return new_para


def main():
    path = os.environ["DOCX_PATH"]
    doc = Document(path)
    paras = doc.paragraphs

    # Chapter 4.1 rewrite
    set_text(
        paras[492],
        "Hệ thống CinemaStar được triển khai theo mô hình phân tách rõ giữa frontend, API gateway và các microservice backend. Cách tổ chức này giúp từng thành phần có thể phát triển, triển khai và bảo trì độc lập, đồng thời vẫn đảm bảo khả năng phối hợp thống nhất thông qua các giao diện API và các kênh giao tiếp nội bộ.",
    )
    set_text(
        paras[494],
        "Mô hình triển khai của hệ thống gồm ba lớp chính: lớp giao diện người dùng, lớp điều phối truy cập và lớp xử lý nghiệp vụ - hạ tầng. Frontend đảm nhiệm phần hiển thị và tương tác; API gateway tiếp nhận, định tuyến và kiểm soát truy cập; các microservice phía sau xử lý từng miền nghiệp vụ riêng biệt như xác thực, người dùng, phim, rạp chiếu, phòng chiếu, ghế, suất chiếu, đặt vé, thanh toán và email.",
    )
    set_text(
        paras[495],
        "Mọi yêu cầu từ giao diện đều đi qua gateway trước khi được chuyển đến service tương ứng. Cách tổ chức này giúp tập trung hóa việc xác thực, phân quyền và định tuyến, đồng thời giữ cho lớp giao diện tách biệt với lớp xử lý nghiệp vụ.",
    )
    set_text(
        paras[496],
        "Việc chia backend thành các service độc lập giúp từng miền chức năng có thể được phát triển và vận hành riêng, hạn chế phụ thuộc chéo và thuận lợi hơn khi mở rộng hoặc thay đổi nghiệp vụ.",
    )
    set_text(
        paras[497],
        "Tầng dữ liệu và hạ tầng bao gồm cơ sở dữ liệu, bộ nhớ đệm và hàng đợi thông điệp. Cơ sở dữ liệu lưu trữ dữ liệu nghiệp vụ chính; cache hỗ trợ truy xuất nhanh các dữ liệu thường dùng; hàng đợi phục vụ các tác vụ bất đồng bộ như gửi email và các công việc nền.",
    )
    set_text(
        paras[498],
        "Nhờ phân lớp như vậy, hệ thống vừa đảm bảo tính nhất quán dữ liệu, vừa cải thiện khả năng phản hồi và khả năng phối hợp giữa các thành phần.",
    )
    set_text(
        paras[499],
        "Tổng thể, mô hình triển khai của CinemaStar đáp ứng được yêu cầu vận hành ổn định, dễ kiểm soát và phù hợp với định hướng phát triển lâu dài của đề tài.",
    )

    set_text(
        paras[501],
        "Quy trình triển khai được tổ chức theo từng bước rõ ràng để bảo đảm hệ thống sau khi đóng gói có thể khởi động và phối hợp ổn định giữa các thành phần.",
    )
    set_text(
        paras[502],
        "Trước hết, mã nguồn của hệ thống được chuẩn bị theo hai khối chính là frontend và backend. Frontend được xây dựng thành ứng dụng web riêng, còn backend được chia thành các service độc lập đi kèm với các thành phần hạ tầng như cơ sở dữ liệu, cache, message queue và lớp điều phối truy cập.",
    )
    set_text(
        paras[503],
        "Sau giai đoạn đóng gói, hệ thống được cấu hình để các lớp có thể giao tiếp với nhau theo đúng vai trò. Frontend kết nối với backend thông qua API gateway, trong khi các service phía sau trao đổi dữ liệu với nhau qua các giao diện nội bộ và các thành phần hạ tầng đã được khai báo trước.",
    )
    set_text(
        paras[504],
        "Khi đưa lên môi trường triển khai, hệ thống được kiểm tra theo từng lớp để bảo đảm trạng thái sẵn sàng vận hành. Các kiểm tra tập trung vào khả năng khởi động của dịch vụ, khả năng truy cập từ giao diện, khả năng định tuyến của gateway và khả năng phối hợp giữa các service phụ thuộc.",
    )
    set_text(
        paras[505],
        "Nhìn chung, quy trình triển khai của CinemaStar được thiết kế theo hướng có kiểm soát và có liên kết rõ ràng giữa giao diện, nghiệp vụ và hạ tầng, nhờ đó việc bảo trì, kiểm thử và mở rộng về sau trở nên thuận lợi hơn.",
    )

    # Chapter 4.2 rewrite
    set_text(
        paras[508],
        "Mục tiêu của quá trình kiểm thử là xác minh tính đúng đắn của các rule nghiệp vụ ở mức đơn vị, đồng thời đánh giá khả năng vận hành của các API chính trên môi trường production. Phạm vi kiểm thử tập trung vào các chức năng quan trọng như xác thực người dùng, tra cứu phim, tra cứu suất chiếu, xem sơ đồ ghế, đánh giá phim, lịch sử đặt vé, thanh toán, báo cáo doanh thu và xuất báo cáo Excel.",
    )
    set_text(
        paras[509],
        "Trong quá trình phát triển, kiểm thử được triển khai theo hai lớp. Ở mức đơn vị, nhóm sử dụng unit test để kiểm tra độc lập các service, controller, scheduler và gRPC nội bộ. Ở mức tích hợp/API, Postman Collection kết hợp Newman CLI được dùng để chạy tự động request, sinh báo cáo HTML/JUnit và đối chiếu kết quả trên môi trường production.",
    )
    set_text(
        paras[510],
        "4.2.2. Công cụ và môi trường kiểm thử",
    )
    set_text(
        paras[511],
        "Ở mức đơn vị, các test chạy trực tiếp trong mã nguồn bằng bộ JUnit/Mockito để kiểm tra từng module. Ở mức API, Postman Collection được sử dụng để xây dựng và quản lý request kiểm thử, trong khi Newman CLI đảm nhiệm việc chạy collection tự động và xuất báo cáo kết quả dưới dạng HTML/JUnit. Các request được gửi trực tiếp đến API gateway production tại địa chỉ https://cinema-api.duckdns.org.",
    )
    set_text(paras[512], "4.2.3. Kiểm thử đơn vị")
    set_text(
        paras[513],
        "Trong quá trình phát triển CinemaStar, nhóm xây dựng bộ unit test cho các lớp xử lý cốt lõi như xác thực, quản lý người dùng, tìm kiếm phim, suất chiếu, đặt vé, thanh toán và đánh giá. Các test này giúp kiểm tra sớm các rule nghiệp vụ quan trọng, hạn chế lỗi khi tích hợp nhiều service và tạo cơ sở tin cậy trước khi triển khai sang môi trường chạy thật.",
    )
    set_text(
        paras[514],
        "Một số unit test tiêu biểu gồm RequestAuthUtilsTest, JwtAuthenticationFilterTest và UserServiceImplTokenFlowTest cho xác thực; CustomerRankServiceImplTest, UserServiceImplPhoneLookupTest và LoyaltyPointsSyncServiceTest cho người dùng; FilmServiceImplTest, FilmControllerSearchIntegrationTest, ActorServiceImplTest và TypeServiceImplTest cho phim; ShowTimeServiceImplTest và SeatSuggestionServiceImplTest cho suất chiếu; BookingServiceImplTest và BookingInternalGrpcServiceTest cho đặt vé; PaymentSessionServiceImplTest, PromotionServiceImplTest và PromotionEngineTest cho thanh toán; cùng ReviewServiceImplTest, UploadServiceImplTest, HallServiceImplTest và CinemaServiceImplTest cho các module nền tảng còn lại.",
    )
    set_text(
        paras[515],
        "Nhìn chung, bộ unit test đóng vai trò lớp kiểm tra sơ bộ trước khi thực hiện kiểm thử tích hợp bằng Newman, giúp phát hiện sớm lỗi logic ngay trong mã nguồn.",
    )

    heading_style = paras[510].style
    body_style = paras[511].style
    api_heading = insert_paragraph_after(
        paras[515],
        "4.2.4. Kiểm thử chức năng API bằng Postman/Newman",
        style=heading_style,
    )
    api_body = insert_paragraph_after(
        api_heading,
        "Postman Collection được sử dụng để nhóm các request theo từng luồng nghiệp vụ, trong khi Newman CLI đảm nhiệm việc chạy tự động collection, thu thập kết quả và xuất báo cáo HTML/JUnit. Cách làm này giúp quá trình kiểm thử API có thể lặp lại nhiều lần trên API gateway production mà không phải thao tác thủ công từng request.",
        style=body_style,
    )
    result_heading = insert_paragraph_after(
        api_body,
        "4.2.5. Kết quả kiểm thử",
        style=heading_style,
    )
    result_body = insert_paragraph_after(
        result_heading,
        "Sau khi chạy collection bằng Newman trên môi trường production, hệ thống thực thi 19 request với tổng cộng 57 assertion. Kết quả cho thấy toàn bộ assertion đều thành công và không có test thất bại.",
        style=body_style,
    )
    insert_paragraph_after(
        result_body,
        "Ngoài kết quả tổng quan, báo cáo Newman còn cung cấp thông tin chi tiết cho từng request. Ví dụ, request POST /api/films/search được sử dụng để kiểm tra chức năng tìm kiếm phim. Kết quả cho thấy API trả về mã trạng thái 200 OK, tỷ lệ kiểm thử thành công đạt 100% và thời gian phản hồi trung bình là 202 ms.",
        style=body_style,
    )

    doc.save(path)

    check = Document(path)
    texts = [p.text.strip() for p in check.paragraphs if p.text.strip()]
    for i in range(490, 520):
        print(f"{i}: {texts[i].encode('unicode_escape').decode('ascii')}")


if __name__ == "__main__":
    main()
