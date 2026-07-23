import os
from docx import Document


def main():
    path = os.environ["DOCX_PATH"]
    doc = Document(path)
    print("PARA", len(doc.paragraphs), "TABLES", len(doc.tables))
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt:
            print(f"P{i}: {txt.encode('unicode_escape').decode('ascii')}")
    for ti, table in enumerate(doc.tables):
        print(f"TABLE {ti} ROWS {len(table.rows)} COLS {len(table.columns)}")
        for ri, row in enumerate(table.rows):
            print(" | ".join(cell.text.encode('unicode_escape').decode('ascii') for cell in row.cells))


if __name__ == "__main__":
    main()
