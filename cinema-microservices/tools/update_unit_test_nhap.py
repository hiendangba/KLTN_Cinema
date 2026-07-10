from copy import deepcopy

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE_DOCX_PATH = r"C:\hoctap\Study\KLTN\File Báo cáo\UnitTest_CinemaStar_Nhap.docx"
OUTPUT_DOCX_PATH = r"C:\hoctap\Study\KLTN\File Báo cáo\UnitTest_CinemaStar_Nhap_capnhat.docx"


ROWS = [
    (
        "Xác thực",
        "3 lớp test / 42 ca kiểm thử",
        "RequestAuthUtilsTest; JwtAuthenticationFilterTest; UserServiceImplTokenFlowTest",
        "Tập trung kiểm tra việc bóc tách thông tin định danh từ request, cơ chế fallback giữa header, cookie và Authorization, khả năng hoạt động của bộ lọc JWT, cùng các luồng login, refresh token, logout và đăng nhập Google. Nhóm test này giúp xác minh các tình huống hợp lệ lẫn các trường hợp token thiếu, sai định dạng hoặc không còn hiệu lực.",
    ),
    (
        "Người dùng",
        "6 lớp test / 32 ca kiểm thử",
        "CustomerRankServiceImplTest; LoyaltyPointsSyncServiceTest; CustomerRankSettlementSyncServiceTest; UserServiceImplPhoneLookupTest; UserServiceImplAuditTest; UserServiceImplLoyaltyPointsTest",
        "Kiểm tra các rule xếp hạng khách hàng, cộng trừ và đồng bộ điểm tích lũy, tra cứu người dùng theo số điện thoại, cũng như các luồng audit và snapshot dữ liệu. Phần này phản ánh rõ cách hệ thống xử lý lịch sử thay đổi dữ liệu người dùng và duy trì tính nhất quán của loyalty workflow.",
    ),
    (
        "Phim",
        "4 lớp test / 17 ca kiểm thử",
        "FilmServiceImplTest; FilmControllerSearchIntegrationTest; ActorServiceImplTest; TypeServiceImplTest",
        "Kiểm tra chức năng tìm kiếm phim, tổng hợp rating summary, ánh xạ dữ liệu diễn viên và thể loại, đồng thời xác minh luồng controller khi nhận request thực tế. Nhóm test này giúp đảm bảo dữ liệu phim trả về đúng cấu trúc và đáp ứng được nhu cầu tra cứu ở phía người dùng.",
    ),
    (
        "Suất chiếu",
        "5 lớp test / 48 ca kiểm thử",
        "ShowTimeServiceImplTest; SeatSuggestionServiceImplTest; PricingPolicyServiceImplTest; ShowTimeStatusSchedulerTest; ShowtimeInternalGrpcServiceTest",
        "Bao phủ các nghiệp vụ tìm kiếm suất chiếu, gợi ý ghế, áp dụng chính sách giá vé theo điều kiện cụ thể, cập nhật trạng thái bằng scheduler và trao đổi dữ liệu nội bộ qua gRPC. Đây là nhóm có mức độ chi tiết cao vì liên quan trực tiếp đến luồng chọn suất chiếu và chuẩn bị dữ liệu trước khi đặt vé.",
    ),
    (
        "Đặt vé",
        "5 lớp test / 51 ca kiểm thử",
        "BookingServiceImplTest; BookingExpirationServiceImplTest; BookingExpirationSchedulerTest; BookingInternalGrpcServiceTest; ProductServiceImplTest",
        "Kiểm tra tạo booking, giữ ghế tạm thời, xử lý hết hạn booking, cập nhật trạng thái ghế trong thời gian thực và các sản phẩm đi kèm như combo. Các test trong nhóm này giúp phát hiện sớm lỗi ở những luồng nghiệp vụ phức tạp, nơi nhiều điều kiện ràng buộc xảy ra đồng thời trước khi giao dịch được xác nhận.",
    ),
    (
        "Thanh toán",
        "6 lớp test / 57 ca kiểm thử",
        "PaymentSessionServiceImplTest; PromotionServiceImplTest; PromotionEngineTest; PaymentInternalGrpcServiceTest; RevenueReportSupportTest; PaymentLoyaltyOutboxServiceTest",
        "Tập trung vào việc tạo và quản lý session thanh toán, áp dụng khuyến mãi, tính toán ưu đãi, trao đổi thanh toán qua gRPC, hỗ trợ báo cáo doanh thu và phát sinh outbox tích điểm. Đây là nhóm có số ca kiểm thử cao nhất, cho thấy phần thanh toán được ưu tiên kiểm soát kỹ do ảnh hưởng trực tiếp đến tiền, ưu đãi và dữ liệu hậu giao dịch.",
    ),
    (
        "Đánh giá & hệ thống",
        "5 lớp test / 44 ca kiểm thử",
        "ReviewServiceImplTest; UploadServiceImplTest; HallServiceImplTest; SeatLayoutServiceImplTest; CinemaServiceImplTest",
        "Kiểm tra điều kiện đánh giá phim, nghiệp vụ upload, quản lý hall, seat layout và các chức năng nền tảng khác của hệ thống rạp. Nhóm này giúp hoàn thiện độ bao phủ cho những module hỗ trợ, bảo đảm hệ thống vẫn ổn định ở các tác vụ quản trị và vận hành phía sau.",
    ),
]


def set_run_font(run, size=11, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def format_cell(cell, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=11):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=11):
    cell.text = text
    format_cell(cell, bold=bold, align=align, size=size)


def insert_paragraph_after(paragraph, text):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p.getparent().remove(new_para._p)
    new_para._p = new_p
    new_para._element = new_p
    new_para.text = text
    return new_para


def main():
    doc = Document(SOURCE_DOCX_PATH)

    doc.paragraphs[2].text = (
        "Tài liệu này tổng hợp 34 lớp unit test tiêu biểu, tương ứng 291 ca kiểm thử "
        "(@Test) thuộc 7 nhóm chức năng chính của CinemaStar. Mục tiêu là giúp bạn trích "
        "nhanh danh sách test, mô tả phạm vi kiểm thử và đưa vào báo cáo theo dạng bảng "
        "có số liệu minh họa rõ ràng hơn."
    )
    for run in doc.paragraphs[2].runs:
        set_run_font(run, size=12)

    heading_para = doc.paragraphs[3]
    old_table = doc.tables[0]
    table_style = old_table.style
    old_tbl = old_table._tbl
    parent = old_tbl.getparent()
    parent.remove(old_tbl)

    new_table = doc.add_table(rows=1, cols=4)
    new_table.style = table_style
    new_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    new_table.autofit = False

    widths = [Cm(2.8), Cm(3.4), Cm(6.2), Cm(6.6)]
    headers = [
        "Nhóm chức năng",
        "Số liệu unit test",
        "Unit test tiêu biểu",
        "Mô tả chi tiết phạm vi kiểm thử",
    ]
    for idx, header in enumerate(headers):
        cell = new_table.rows[0].cells[idx]
        cell.width = widths[idx]
        set_cell_text(cell, header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)

    for group, stats, tests, desc in ROWS:
        row = new_table.add_row()
        values = [group, stats, tests, desc]
        aligns = [
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.JUSTIFY,
        ]
        for idx, value in enumerate(values):
            row.cells[idx].width = widths[idx]
            set_cell_text(row.cells[idx], value, align=aligns[idx], size=11)

    new_tbl = new_table._tbl
    parent = new_tbl.getparent()
    parent.remove(new_tbl)
    heading_para._p.addnext(new_tbl)

    doc.paragraphs[8].text = (
        "Nếu cần viết ngắn gọn trong Chương 4, bạn có thể sử dụng trực tiếp cột số liệu unit "
        "test và cột mô tả chi tiết trong bảng trên, sau đó giữ lại 2 đoạn mẫu ở bên dưới "
        "làm phần thuyết minh."
    )
    for run in doc.paragraphs[8].runs:
        set_run_font(run, size=12)

    doc.save(OUTPUT_DOCX_PATH)


if __name__ == "__main__":
    main()
